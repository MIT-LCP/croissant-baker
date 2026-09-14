"""Spreadsheet handler: one record set per sheet that holds a single table.

A sheet holding anything else is reported rather than guessed at. What counts
as a table, and every reason one is refused, is in
``docs/user-guide/supported-formats.md``.
"""

import itertools
import logging
import re
import shutil
import tempfile
import zipfile
from dataclasses import dataclass
from typing import Iterator, List, Optional, Tuple

import mlcroissant as mlc
import openpyxl
import xlrd
from openpyxl.packaging.custom import CustomPropertyList
from openpyxl.xml.functions import fromstring

from croissant_baker.entries import Diagnostic, DiagnosticCode
from croissant_baker.handlers.base_handler import BuildResult, FileTypeHandler
from croissant_baker.handlers.utils import (
    SCHEMA_SAMPLE,
    allocate_record_set_ids,
    display_name,
    infer_croissant_type,
    make_field_id,
)
from croissant_baker.sources import FileSource

logger = logging.getLogger(__name__)

_XLSX = ".xlsx"
_XLSM = ".xlsm"
_XLS = ".xls"

# Claimed on the signature, not the name: lab systems export HTML and TSV
# under a spreadsheet suffix. OLE2 under a modern suffix is an encrypted
# workbook, and "no handler" is the wrong thing to tell someone holding one.
_ZIP_MAGIC = b"PK\x03\x04"
_OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"

_MAGIC = {
    _XLSX: (_ZIP_MAGIC, _OLE2_MAGIC),
    _XLSM: (_ZIP_MAGIC, _OLE2_MAGIC),
    _XLS: (_OLE2_MAGIC,),
}

_MEDIA_TYPES = {
    _XLSX: "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    _XLSM: "application/vnd.ms-excel.sheet.macroEnabled.12",
    _XLS: "application/vnd.ms-excel",
}

#: How far down a sheet a header may sit, bounded the way the CSV preamble is.
_HEADER_SEARCH_ROWS = 100

#: How many worksheet rows are inspected per sheet. Bounds traversal only:
#: hashing, decompression and the shared string table are outside it.
_SCAN_ROWS = 1000

_NUMERIC_TEXT = re.compile(r"^[+-]?(\d+\.?\d*|\.\d+)([eE][+-]?\d+)?$")

_SPOOL_MAX = 8 * 1024 * 1024
_COPY_CHUNK = 64 * 1024
_CUSTOM_PART = "docProps/custom.xml"
_CONTENT_TYPE_ENTRY = re.compile(
    rb'<Override[^>]*PartName="/?docProps/custom\.xml"[^>]*/>'
)
_RELATIONSHIP_ENTRY = re.compile(
    rb'<Relationship[^>]*Target="/?docProps/custom\.xml"[^>]*/>'
)


class _Spooled(tempfile.SpooledTemporaryFile):
    """A spooled temporary file that ``zipfile`` will accept on Python 3.10.

    ``seekable``/``readable``/``writable`` reached it only in 3.11
    (python/cpython#70363), and ``zipfile`` asks for them.
    """

    def seekable(self) -> bool:
        return True

    def readable(self) -> bool:
        return True

    def writable(self) -> bool:
        return True


@dataclass(frozen=True)
class _Sheet:
    """A table, a reason there is none, or neither for an empty sheet."""

    table: Optional[dict] = None
    skipped: str = ""


def _blank(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _positions(row: tuple) -> frozenset:
    """The 1-based column positions of ``row`` that hold something."""
    return frozenset(i for i, value in enumerate(row, 1) if not _blank(value))


def _is_label(value) -> bool:
    """Whether a header cell reads as an authored name rather than a value."""
    return isinstance(value, str) and not _NUMERIC_TEXT.match(value.strip())


def _labelled(row: tuple, positions: frozenset) -> bool:
    """Whether a strict majority of a header candidate's cells are labels."""
    labels = sum(1 for p in positions if _is_label(row[p - 1]))
    return labels * 2 > len(positions)


def _column_type(values: list) -> str:
    """The Croissant type for one column, by majority of its sampled values."""
    if not values:
        return "sc:Text"
    votes: dict = {}
    for value in values:
        name = infer_croissant_type(value)
        votes[name] = votes.get(name, 0) + 1
    winner = max(votes, key=votes.get)
    if winner == "cr:Int64" and "cr:Float64" in votes:
        return "cr:Float64"
    return winner


def _scan(rows: Iterator[tuple]) -> Tuple[List[tuple], bool]:
    """``(the inspected rows, whether the sheet ended inside the budget)``.

    One row of lookahead, so a table filling the window exactly is not
    reported as a lower bound.
    """
    try:
        window = list(itertools.islice(rows, _SCAN_ROWS))
        ended = next(rows, None) is None
    finally:
        close = getattr(rows, "close", None)
        if close is not None:
            close()
    return window, ended


def _regions(window: List[tuple], filled: List[frozenset]) -> List[List[int]]:
    """The row indices of each table region, blank rows between them.

    A lone row of labels above a region of the same width joins it: a gap
    under a header is a spacer, not the end of the table.
    """
    runs, start = [], None
    for index, positions in enumerate(filled):
        if positions and start is None:
            start = index
        elif not positions and start is not None:
            runs.append(list(range(start, index)))
            start = None
    if start is not None:
        runs.append(list(range(start, len(filled))))

    out: List[List[int]] = []
    for run in runs:
        head = out[-1] if out else None
        if (
            head is not None
            and len(head) == 1
            and max(filled[head[0]])
            == max(frozenset().union(*(filled[i] for i in run)))
            and _labelled(window[head[0]], filled[head[0]])
        ):
            out[-1] = head + run
        else:
            out.append(run)
    return out


def _all_labels(row: tuple, positions: frozenset) -> bool:
    return all(_is_label(row[p - 1]) for p in positions)


def _find_header(
    window: List[tuple], filled: List[frozenset], rows: List[int], width: int
) -> Tuple[Optional[int], str]:
    """``(the header's offset within ``rows``, why there is none)``.

    The region's own first row is asked first, so a data row cannot outrank a
    header that merely has a gap in it.
    """
    span = frozenset(range(1, width + 1))
    head = filled[rows[0]]

    if max(head) == width:
        if not _labelled(window[rows[0]], head):
            return None, "its first row reads as data rather than column names"
        if head == span:
            return 0, ""
        if len(rows) < 2:
            return None, "it holds a single row, with no data under it"
        below = filled[rows[1]]
        if not (span - head) <= below:
            return None, "no row in it names every column"
        if below != span and _all_labels(window[rows[1]], below):
            return None, "its first two rows both read as column names"
        return 0, ""

    if len(rows) > 1:
        below = filled[rows[1]]
        if (
            below != span
            and head | below >= span
            and _labelled(window[rows[0]], head)
            and _labelled(window[rows[1]], below)
        ):
            return None, "its first two rows both read as column names"

    spanning = next((n for n, i in enumerate(rows) if filled[i] >= span), None)
    if spanning is None or not _all_labels(
        window[rows[spanning]], filled[rows[spanning]]
    ):
        return None, (
            "a row in it holds more columns than the row above it names, so it "
            "reads as two tables rather than one"
        )
    return spanning, ""


def _read_region(
    window: List[tuple], filled: List[frozenset], rows: List[int]
) -> Tuple[Optional[dict], str]:
    """``(the table this region holds, why it holds none)``."""
    populated = frozenset().union(*(filled[i] for i in rows))
    width = max(populated)
    missing = frozenset(range(1, width + 1)) - populated
    if missing:
        return None, (
            f"column {min(missing)} is empty throughout, so it reads as two "
            "tables side by side rather than one"
        )

    header_at, why = _find_header(window, filled, rows, width)
    if header_at is None:
        return None, why
    if header_at + 1 >= len(rows):
        return None, "its header has no data under it"
    if header_at and _read_region(window, filled, rows[:header_at])[0] is not None:
        return None, (
            "a row in it holds more columns than the row above it names, so "
            "it reads as two tables rather than one"
        )

    head = rows[header_at]
    header = window[head]
    columns = [
        str(header[p - 1]).strip() if p in filled[head] else f"column_{p}"
        for p in range(1, width + 1)
    ]

    samples: list = [[] for _ in columns]
    num_rows = 0
    for row in (window[i] for i in rows[header_at + 1 :]):
        num_rows += 1
        if num_rows > SCHEMA_SAMPLE:
            continue
        for index in range(width):
            value = row[index] if index < len(row) else None
            if not _blank(value):
                samples[index].append(value)

    return {
        "columns": columns,
        "column_types": [_column_type(values) for values in samples],
        "num_rows": num_rows,
    }, ""


def _read_sheet(rows: Iterator[tuple]) -> _Sheet:
    """Read one worksheet as the single table it is supposed to hold."""
    window, ended = _scan(rows)
    filled = [_positions(row) for row in window]
    if not any(filled):
        return _Sheet()

    tables, refusals, late = [], [], 0
    for rows in _regions(window, filled):
        table, why = _read_region(window, filled, rows)
        # A header is only looked for near the top of a sheet, but a table
        # below that still makes the sheet two tables, and saying so is what
        # the limit above is for.
        if table is None:
            if rows[0] < _HEADER_SEARCH_ROWS:
                refusals.append(why)
            continue
        if rows[0] >= _HEADER_SEARCH_ROWS:
            late += 1
            continue
        table["row_count_is_exact"] = ended or rows[-1] < len(window) - 1
        tables.append(table)

    if len(tables) + late > 1:
        return _Sheet(
            skipped=(
                f"it holds {len(tables) + late} tables in its first {_SCAN_ROWS} "
                "rows, and which one describes the sheet would be a guess"
            )
        )
    if tables:
        return _Sheet(table=tables[0])
    if refusals:
        return _Sheet(skipped=refusals[0])
    return _Sheet(
        skipped=(
            f"no table starts in its first {_HEADER_SEARCH_ROWS} rows, which is "
            "as far down a sheet a header is looked for"
        )
    )


def _xls_value(cell, datemode: int):
    """One legacy cell as the object the modern reader would have returned.

    BIFF stores numbers as doubles, dates as serial offsets and errors as
    codes, so without this the same table types differently in each format.
    """
    if cell.ctype == xlrd.XL_CELL_DATE:
        return xlrd.xldate.xldate_as_datetime(cell.value, datemode)
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return bool(cell.value)
    if cell.ctype == xlrd.XL_CELL_ERROR:
        return xlrd.error_text_from_code.get(cell.value)
    if cell.ctype == xlrd.XL_CELL_NUMBER:
        return int(cell.value) if float(cell.value).is_integer() else cell.value
    if cell.ctype == xlrd.XL_CELL_TEXT:
        return cell.value
    return None


def _custom_properties_are_broken(stream) -> bool:
    """Whether this workbook's optional custom-properties part will not parse.

    Asked only after a read has failed, to license dropping that one part. If
    it parses, the failure was something else and the workbook stays broken.
    """
    try:
        stream.seek(0)
        with zipfile.ZipFile(stream) as archive:
            if _CUSTOM_PART not in archive.namelist():
                return False
            tree = fromstring(archive.read(_CUSTOM_PART))
        CustomPropertyList.from_tree(tree)
    except zipfile.BadZipFile:
        return False
    except Exception:
        return True
    return False


def _edited(item: zipfile.ZipInfo) -> Optional[re.Pattern]:
    """The entry to strike out of this package part, if it names one."""
    if item.filename == "[Content_Types].xml":
        return _CONTENT_TYPE_ENTRY
    if item.filename.endswith("_rels/.rels"):
        return _RELATIONSHIP_ENTRY
    return None


def _without_custom_properties(stream):
    """A copy of the workbook with the custom-properties part left out.

    Copied between open handles in bounded chunks, so a worksheet is never
    held whole. The original file is untouched.
    """
    stream.seek(0)
    copy = _Spooled(max_size=_SPOOL_MAX)
    try:
        with (
            zipfile.ZipFile(stream) as archive,
            zipfile.ZipFile(copy, "w", zipfile.ZIP_DEFLATED) as out,
        ):
            for item in archive.infolist():
                if item.filename == _CUSTOM_PART:
                    continue
                pattern = _edited(item)
                if pattern is not None:
                    out.writestr(item, pattern.sub(b"", archive.read(item.filename)))
                    continue
                with archive.open(item) as src, out.open(item, "w") as dst:
                    shutil.copyfileobj(src, dst, _COPY_CHUNK)
    except Exception:
        copy.close()
        raise
    copy.seek(0)
    return copy


def _sheets_of(book) -> List[tuple]:
    """Every worksheet of an OOXML workbook, one at a time.

    A sheet that fails while its rows are read costs that sheet. Anything
    shared raises out of here, because without it no sheet is trustworthy.
    """
    out = []
    for sheet in book.worksheets:
        try:
            # read_only trusts the declared dimensions, and a writer that
            # understates them silently drops rows and columns. Resetting makes
            # openpyxl infer the extent from the cells it actually sees.
            # https://openpyxl.readthedocs.io/en/stable/optimized.html
            sheet.reset_dimensions()
            read = _read_sheet(sheet.iter_rows(values_only=True))
        except Exception as exc:  # noqa: BLE001 — one sheet, not the workbook
            read = _Sheet(skipped=f"it could not be read: {exc}")
        out.append((sheet.title, getattr(sheet, "sheet_state", "visible"), read))
    return out


def _read_with_openpyxl(source: FileSource) -> Tuple[List[tuple], bool]:
    """Read OOXML, streaming each sheet and closing the archive afterwards."""
    with source.open() as stream:
        recovered = None
        try:
            book = openpyxl.load_workbook(stream, read_only=True, data_only=True)
        except Exception:
            if not _custom_properties_are_broken(stream):
                raise
            recovered = _without_custom_properties(stream)
            try:
                book = openpyxl.load_workbook(recovered, read_only=True, data_only=True)
            except Exception:
                recovered.close()
                raise
        try:
            return _sheets_of(book), recovered is not None
        finally:
            book.close()
            if recovered is not None:
                recovered.close()


_XLS_VISIBILITY = {0: "visible", 1: "hidden", 2: "veryHidden"}


def _read_with_xlrd(source: FileSource) -> Tuple[List[tuple], bool]:
    """Read the legacy format, which has no streaming reader.

    The whole file is held in memory, bounded by the format's own 65,536-row
    ceiling. The row budget applies here too, so both readers look as far.
    """
    with source.open() as stream:
        book = xlrd.open_workbook(file_contents=stream.read())
    try:
        out = []
        for sheet in book.sheets():
            try:
                read = _read_sheet(
                    tuple(_xls_value(cell, book.datemode) for cell in sheet.row(index))
                    for index in range(sheet.nrows)
                )
            except Exception as exc:  # noqa: BLE001 — one sheet, not the workbook
                read = _Sheet(skipped=f"it could not be read: {exc}")
            visibility = _XLS_VISIBILITY.get(getattr(sheet, "visibility", 0), "visible")
            out.append((sheet.name, visibility, read))
        return out, False
    finally:
        book.release_resources()


def _rows_phrase(sheet: dict) -> str:
    """How many rows a sheet holds, or how many it is known to hold at least."""
    if sheet["row_count_is_exact"]:
        return f"{sheet['num_rows']} rows"
    return f"at least {sheet['num_rows']} rows"


class SpreadsheetHandler(FileTypeHandler):
    """Handler for Excel workbooks, one record set per readable sheet."""

    EXTENSIONS = (_XLS, _XLSM, _XLSX)
    FORMAT_NAME = "Spreadsheet"
    FORMAT_DESCRIPTION = "Per sheet: column names, inferred types, row count"

    def claims(self, source: FileSource) -> bool:
        magic = _MAGIC.get(source.suffix)
        if magic is None or not source.exists:
            return False
        head = source.peek(max(len(m) for m in magic))
        if any(head.startswith(m) for m in magic):
            return True
        # Debug, not warning: the generator names every undescribed file once
        # through a capped path, so warning here doubles it and escapes the cap.
        logger.debug(
            "Skipping %s: it does not open with the signature its name implies",
            source.relative_path,
        )
        return False

    def extract(self, source: FileSource, **kwargs) -> dict:
        """Describe every sheet of one workbook that holds a single table."""
        if not source.exists:
            raise FileNotFoundError(f"Spreadsheet not found: {source.relative_path}")

        if source.suffix != _XLS and source.peek(len(_OLE2_MAGIC)) == _OLE2_MAGIC:
            raise ValueError(
                f"{source.relative_path} is named for the modern format but is an "
                "OLE2 container, which means it is either encrypted or a legacy "
                "workbook under the wrong extension. Supply an unencrypted "
                "workbook, saved with the extension its format calls for."
            )

        reader = _read_with_xlrd if source.suffix == _XLS else _read_with_openpyxl
        try:
            sheets, ignored_properties = reader(source)
        except Exception as exc:
            raise ValueError(
                f"Failed to read spreadsheet {source.relative_path}: {exc}"
            ) from exc

        described, diagnostics = [], []
        for name, visibility, read in sheets:
            if read.table is not None:
                described.append({"name": name, "visibility": visibility, **read.table})
            elif read.skipped:
                diagnostics.append(
                    Diagnostic(DiagnosticCode.SHEET_SKIPPED, read.skipped, part=name)
                )

        if ignored_properties:
            diagnostics.append(
                Diagnostic(
                    DiagnosticCode.PROPERTIES_IGNORED,
                    "its custom document properties are malformed and were "
                    "ignored; nothing the document carries comes from them",
                )
            )

        if not described:
            raise ValueError(
                f"No sheet of {source.relative_path} holds a single table: "
                + "; ".join(f"{d.part!r} {d.detail}" for d in diagnostics)
            )

        return {
            "file_name": source.name,
            "file_size": source.size,
            "sha256": source.sha256,
            "encoding_format": _MEDIA_TYPES[source.suffix],
            "sheets": described,
            "diagnostics": diagnostics,
            "description": _file_description(diagnostics),
        }

    def build_croissant(self, file_metas: list, file_ids: list) -> tuple:
        """One RecordSet per sheet, all sourced from the workbook's FileObject."""
        # Sorted: allocation order decides which candidate is moved, and
        # batch order is discovery order.
        suffixes = sorted(
            {sheet["name"] for meta in file_metas for sheet in meta["sheets"]}
        )
        allocated = allocate_record_set_ids(file_metas, suffixes)

        record_sets: list = []
        for file_id, meta, ids in zip(file_ids, file_metas, allocated):
            shown = display_name(meta)
            for sheet in meta["sheets"]:
                record_sets.append(
                    self._record_set(sheet, file_id, ids[sheet["name"]], shown)
                )
        return BuildResult([], record_sets)

    @staticmethod
    def _record_set(
        sheet: dict, file_id: str, record_set_id: str, shown: str
    ) -> mlc.RecordSet:
        used_field_ids: set = set()
        fields = [
            mlc.Field(
                id=make_field_id(record_set_id, column, used_field_ids),
                name=column,
                description=f"Column '{column}' of sheet '{sheet['name']}' in {shown}",
                data_types=[column_type],
                source=mlc.Source(file_object=file_id),
            )
            for column, column_type in zip(sheet["columns"], sheet["column_types"])
        ]
        hidden = "" if sheet["visibility"] == "visible" else f", {sheet['visibility']}"
        return mlc.RecordSet(
            id=record_set_id,
            name=sheet["name"],
            description=(
                f"Sheet '{sheet['name']}' of {shown} ({_rows_phrase(sheet)}{hidden})"
            ),
            fields=fields,
        )


def _file_description(diagnostics: list) -> Optional[str]:
    """What the manifest says about the parts of a workbook it does not carry.

    On the FileObject, because a document outlives the run that made it.
    """
    if not diagnostics:
        return None
    return " ".join(
        f"Sheet {d.part!r} was not described because {d.detail}."
        if d.part
        else f"{d.detail[0].upper()}{d.detail[1:]}."
        for d in diagnostics
    )
