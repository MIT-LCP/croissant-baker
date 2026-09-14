"""What the shared contract sweep cannot express about workbooks."""

import datetime
import hashlib
import io
import json
from pathlib import Path

import openpyxl
import pytest

import croissant_baker.handlers.spreadsheet_handler as spreadsheet_handler_module
from croissant_baker.entries import DiagnosticCode
from croissant_baker.handlers.spreadsheet_handler import SpreadsheetHandler
from croissant_baker.sources import make_source
from tests.helpers import (
    DATA,
    bake,
    bake_with_report,
    cli,
    file_objects,
    record_sets,
)

_LOGGER = "croissant_baker.handlers.spreadsheet_handler"
_LEGACY = DATA / "spreadsheets" / "manifest.xls"

#: The rows of the committed .xls fixture, to be written as its modern twin.
_MANIFEST = [
    ["sample_id", "tissue", "rin", "reads", "passed", "collected"],
    ["S1", "lung", 9.1, 400000000, True, datetime.date(2024, 3, 1)],
    ["S2", "liver", 8.4, 120000000, False, datetime.date(2024, 3, 2)],
    # A row with nothing on it but an error, which BIFF stores as a code. The
    # modern reader hands back the same text for a real error cell.
    [None, None, None, None, None, "#N/A"],
]

_SIDE_BY_SIDE = [
    ["metric", "value", None, "note", "author"],
    ["depth", 30, None, "QC passed", "jd"],
]

#: One sheet layout each, and what it reads as: a ``(columns, rows)`` pair when
#: it describes, the reason it is refused for when it does not, or ``None`` when
#: it is passed over without a word.
_LAYOUTS = {
    "preamble-above-the-header": (
        [
            ["Supplementary Table 1. Differentially expressed genes"],
            [],
            ["Contact", "j.doe@lab.example"],
            ["gene", "log2FC", "padj", "tissue"],
            ["TP53", -2.1, 0.001, "lung"],
        ],
        (["gene", "log2FC", "padj", "tissue"], 1),
    ),
    "cells-the-writer-left-empty": (
        [["gene", "padj", None], ["TP53", 0.001, None]],
        (["gene", "padj"], 1),
    ),
    "a-blank-row-ends-the-table": (
        [["gene", "padj"], ["TP53", 0.001], [], ["*p<0.05, two-sided"]],
        (["gene", "padj"], 1),
    ),
    "whitespace-counts-as-blank": (
        [[" gene ", "padj"], ["TP53", 0.001], ["   ", "  "], ["MYC", 0.02]],
        (["gene", "padj"], 1),
    ),
    "a-gap-under-the-header": (
        [["gene", "padj"], [], ["TP53", 0.001], ["MYC", 0.02]],
        (["gene", "padj"], 2),
    ),
    "more-rows-than-the-header-window": (
        [["gene", "padj"]] + [[f"G{n}", 0.01] for n in range(150)],
        (["gene", "padj"], 150),
    ),
    # A pandas export: to_excel leaves A1 empty over the index. The row under
    # it establishes the column, so the column is named for its position and
    # the first data row stays data.
    "a-blank-label-over-an-index": (
        [[None, "gene", "expr", "pval"], [0, "TP53", 1.2, 0.01], [1, "MYC", 0.4, 0.2]],
        (["column_1", "gene", "expr", "pval"], 2),
    ),
    # A text index column, which the numeric case above cannot stand in for:
    # its data row reads as labels too, so the first row has to be asked first.
    "a-blank-label-over-a-text-index": (
        [[None, "gene", "tissue"], ["S1", "TP53", "lung"], ["S2", "MYC", "liver"]],
        (["column_1", "gene", "tissue"], 2),
    ),
    # A merged upper layer whose lower row names every column: the leaf names
    # are taken and the grouping is lost, because nothing in a read-only
    # worksheet separates this from a contact block above a header.
    "a-merged-header-whose-lower-row-is-complete": (
        [["Patient", "BP", None], ["ID", "systolic", "diastolic"], ["P1", 120, 80]],
        (["ID", "systolic", "diastolic"], 1),
    ),
    "a-second-table-past-the-header-window": (
        [["gene", "padj"]]
        + [[f"G{n}", 0.01] for n in range(110)]
        + [[], ["a", "b"], ["1", 2]],
        "holds 2 tables",
    ),
    "a-blank-label-inside-the-header": (
        [["gene", None, "pval"], ["TP53", "lung", 0.01]],
        (["gene", "column_2", "pval"], 1),
    ),
    "a-merged-header-over-two-rows": (
        [["Patient", "BP", None], [None, "systolic", "diastolic"], ["P1", 120, 80]],
        "both read as column names",
    ),
    "a-numeric-header": ([[1, 2, 3], [4, 5, 6]], "reads as data rather than"),
    "a-date-header": (
        [[datetime.date(2024, 1, 1), datetime.date(2024, 1, 2)], [1, 2]],
        "reads as data rather than",
    ),
    "numbers-the-writer-stored-as-text": (
        [["2023", "2024"], [1, 2]],
        "reads as data rather than",
    ),
    "duplicate-labels": ([["value", "value"], [1, "text"]], (["value", "value"], 1)),
    "two-tables-stacked-the-second-wider": (
        [["gene", "padj"], ["TP53", 0.01], [], ["a", "b", "c"], ["1", "2", "3"]],
        "holds 2 tables",
    ),
    "two-tables-stacked-the-same-width": (
        [["gene", "padj"], ["TP53", 0.01], [], ["gene", "padj"], ["MYC", 0.02]],
        "holds 2 tables",
    ),
    "a-later-row-wider-than-the-header": (
        [["gene", "padj"], ["TP53", 0.01], ["MYC", 0.02, "x"], ["A", 1.0, "y"]],
        "more columns than the row above it names",
    ),
    "two-tables-side-by-side": (_SIDE_BY_SIDE, "column 3 is empty throughout"),
    "header-with-no-data": ([["gene", "log2FC"]], "no data under it"),
    "a-table-below-the-header-window": (
        [[] for _ in range(120)] + [["gene", "padj"], ["TP53", 0.01]],
        "no table starts in its first 100 rows",
    ),
    "empty": ([], None),
}

#: One column's values each, and the Croissant type they are read as. The
#: subject column rides with a second column so no row is wholly blank.
_TYPES = {
    "a-fraction-widens-the-column": ([1, 2, 2.5], "cr:Float64"),
    "the-odd-value-out-loses": ([1, 2, 3, "n/a"], "cr:Int64"),
    "blank-cells-are-not-values": ([1, None, None], "cr:Int64"),
    "nothing-at-all": ([None, None], "sc:Text"),
    "dates": (
        [datetime.datetime(2024, 3, 1), datetime.datetime(2024, 3, 2)],
        "sc:DateTime",
    ),
    "times": ([datetime.time(9, 30), datetime.time(10, 0)], "sc:Time"),
    "formulas-with-no-cached-value-are-blank": (["=1+1", "=2+2", 5], "cr:Int64"),
}


@pytest.fixture
def handler() -> SpreadsheetHandler:
    return SpreadsheetHandler()


def _book(path: Path, sheets: dict) -> Path:
    book = openpyxl.Workbook()
    book.remove(book.active)
    for title, rows in sheets.items():
        sheet = book.create_sheet(title)
        for row in rows:
            sheet.append(row)
    book.save(path)
    return path


def _understate_dimensions(path: Path) -> None:
    """Rewrite a worksheet's declared extent to less than it holds."""
    import re
    import zipfile

    source = path.read_bytes()
    with zipfile.ZipFile(io.BytesIO(source)) as archive:
        items = [(i, archive.read(i.filename)) for i in archive.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as out:
        for item, data in items:
            if item.filename == "xl/worksheets/sheet1.xml":
                data = re.sub(rb"<dimension[^>]*>", b'<dimension ref="A1:A2"/>', data)
            out.writestr(item, data)


def _sheets(handler: SpreadsheetHandler, path: Path) -> list:
    return handler.extract(make_source(path))["sheets"]


@pytest.mark.parametrize(("rows", "expected"), _LAYOUTS.values(), ids=_LAYOUTS)
def test_a_sheet_is_read_as_the_one_table_it_holds(
    handler: SpreadsheetHandler, tmp_path: Path, rows: list, expected
) -> None:
    """Each layout rides with a sheet that does read, so a refusal shows up as
    an absence rather than as a failed workbook."""
    path = _book(tmp_path / "layout.xlsx", {"anchor": [["a"], [1]], "subject": rows})
    meta = handler.extract(make_source(path))

    described = {sheet["name"]: sheet for sheet in meta["sheets"]}
    assert "anchor" in described, "the anchor sheet stopped reading"
    reasons = {d.part: d.detail for d in meta["diagnostics"]}

    if isinstance(expected, tuple):
        sheet = described["subject"]
        assert (sheet["columns"], sheet["num_rows"]) == expected
        assert "subject" not in reasons
    elif expected is None:
        assert "subject" not in described
        assert "subject" not in reasons, "an empty sheet should be passed over"
    else:
        assert "subject" not in described
        assert expected in reasons["subject"]


@pytest.mark.parametrize(("values", "expected"), _TYPES.values(), ids=_TYPES)
def test_a_column_is_typed_by_the_values_it_holds(
    handler: SpreadsheetHandler, tmp_path: Path, values: list, expected: str
) -> None:
    """A fraction anywhere widens a whole-number column, which is the
    promotion the CSV reader makes for the same case."""
    rows = [["subject", "anchor"]] + [[value, "x"] for value in values]
    path = _book(tmp_path / "types.xlsx", {"s": rows})

    (sheet,) = _sheets(handler, path)

    assert sheet["column_types"][0] == expected


def test_each_sheet_becomes_a_record_set_of_its_own_columns(dataset: Path) -> None:
    """Asserted on the baked document, because ``no extract`` is only visible
    there: mlc.Source fills in an empty Extract that never reaches the JSON."""
    _book(
        dataset / "study.xlsx",
        {
            "samples": _MANIFEST,
            "platforms": [["platform_id", "reads"], ["GPL24676", 400000000]],
        },
    )

    sets = {node["name"]: node for node in record_sets(bake(dataset))}

    assert {
        name: [f["name"] for f in node["field"]] for name, node in sets.items()
    } == {
        "samples": ["sample_id", "tissue", "rin", "reads", "passed", "collected"],
        "platforms": ["platform_id", "reads"],
    }
    # dataType round-trips as an rdflib URIRef, whose __eq__ refuses a str.
    assert [str(f["dataType"]) for f in sets["samples"]["field"]] == [
        "sc:Text",
        "sc:Text",
        "cr:Float64",
        "cr:Int64",
        "sc:Boolean",
        "sc:DateTime",
    ]
    assert all(
        field["source"] == {"fileObject": {"@id": "file_0"}}
        for node in sets.values()
        for field in node["field"]
    )


def test_a_workbook_holding_no_table_fails_by_name(
    handler: SpreadsheetHandler, tmp_path: Path
) -> None:
    """The other branch: a workbook that reads but says nothing is a failure
    with a reason, not an empty description."""
    path = _book(tmp_path / "layout.xlsx", {"grid": _SIDE_BY_SIDE})

    with pytest.raises(ValueError, match="layout.xlsx"):
        handler.extract(make_source(path))


def test_a_legacy_workbook_describes_like_its_modern_twin(
    handler: SpreadsheetHandler, tmp_path: Path
) -> None:
    """BIFF stores whole numbers and booleans as doubles and dates as serial
    offsets, so the two readers agree only because the .xls adapter converts
    all three back. The fixture carries one column of each."""
    modern = _book(tmp_path / "manifest.xlsx", {"manifest": _MANIFEST})

    legacy = _sheets(handler, _LEGACY)

    assert [(sheet["columns"], sheet["column_types"]) for sheet in legacy] == [
        (
            ["sample_id", "tissue", "rin", "reads", "passed", "collected"],
            [
                "sc:Text",
                "sc:Text",
                "cr:Float64",
                "cr:Int64",
                "sc:Boolean",
                "sc:DateTime",
            ],
        )
    ]
    assert legacy == _sheets(handler, modern)


def test_each_workbook_format_reports_its_own_media_type(
    handler: SpreadsheetHandler, tmp_path: Path
) -> None:
    """A macro-enabled workbook is an OOXML zip, so it reads like .xlsx and is
    typed unlike it. Nothing else in the suite carries one."""
    table = {"s": [["a"], [1]]}

    assert [
        handler.extract(make_source(path))["encoding_format"]
        for path in (
            _book(tmp_path / "book.xlsx", table),
            _book(tmp_path / "book.xlsm", table),
            _LEGACY,
        )
    ] == [
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/vnd.ms-excel.sheet.macroEnabled.12",
        "application/vnd.ms-excel",
    ]


@pytest.mark.parametrize(
    ("name", "payload"),
    [
        ("~$study.xlsx", b"\x20\x00j\x00d\x00\x00\x00"),
        ("export.xls", b"sample_id\ttissue\nS1\tlung\n"),
    ],
    ids=["lock-file", "tsv-named-xls"],
)
def test_a_file_that_is_not_a_workbook_is_declined_at_debug(
    handler: SpreadsheetHandler,
    tmp_path: Path,
    caplog: pytest.LogCaptureFixture,
    name: str,
    payload: bytes,
) -> None:
    """One case per signature. Lab systems export TSV and HTML under a
    spreadsheet name, and Excel's ~$ lock files begin with no dot, so discovery
    offers both here."""
    path = tmp_path / name
    path.write_bytes(payload)

    with caplog.at_level("DEBUG", logger=_LOGGER):
        assert handler.claims(make_source(path)) is False

    assert [
        r for r in caplog.records if r.levelname == "DEBUG" and name in r.message
    ], caplog.records


def test_both_readers_are_released_after_extraction(
    handler: SpreadsheetHandler, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """read_only mode holds the archive open and xlrd holds the whole file, so
    a bake of a directory of workbooks must not accumulate either. Nothing but
    the release call distinguishes a closed reader from a collected one."""
    released = []

    def watch(module, opener: str, release: str) -> None:
        original = getattr(module, opener)

        def spy(*args, **kwargs):
            book = original(*args, **kwargs)
            close = getattr(book, release)

            def watched():
                released.append(book)
                return close()

            setattr(book, release, watched)
            return book

        monkeypatch.setattr(module, opener, spy)

    watch(spreadsheet_handler_module.openpyxl, "load_workbook", "close")
    watch(spreadsheet_handler_module.xlrd, "open_workbook", "release_resources")

    handler.extract(make_source(_book(tmp_path / "b.xlsx", {"s": [["a"], [1]]})))
    handler.extract(make_source(_LEGACY))

    assert len(released) == 2, f"a reader was left open: {released}"


def test_names_that_sanitize_alike_still_get_distinct_ids(
    handler: SpreadsheetHandler, tmp_path: Path
) -> None:
    """@id is unique document-wide, so a collision raises during assembly
    rather than quietly dropping a sheet or a column. Two columns of a sheet
    may share a header, and each keeps the type its own values earned."""
    path = _book(
        tmp_path / "plates.xlsx",
        {"A B": [["value", "value"], [1, "text"]], "A+B": [["y"], [2]]},
    )
    meta = handler.extract(make_source(path))
    meta["relative_path"] = "plates.xlsx"

    built = handler.build_croissant([meta], ["file_0"])

    assert [node.id for node in built.record_sets] == [
        "plates_A_B",
        "plates_A_B__2",
    ]
    assert [
        (field.id, str(field.data_types[0])) for field in built.record_sets[0].fields
    ] == [
        ("plates_A_B/value", "cr:Int64"),
        ("plates_A_B/value__2", "sc:Text"),
    ]


class _CountingRows:
    """Records how far it was walked and whether it closed."""

    def __init__(self, total: int) -> None:
        self.total, self.seen, self.closed = total, 0, False

    def __iter__(self) -> "_CountingRows":
        return self

    def __next__(self) -> tuple:
        if self.seen >= self.total:
            raise StopIteration
        self.seen += 1
        return ("gene", "padj") if self.seen == 1 else (f"G{self.seen}", 0.01)

    def close(self) -> None:
        self.closed = True


_BUDGET = spreadsheet_handler_module._SCAN_ROWS


@pytest.mark.parametrize(
    ("total", "rows", "exact"),
    [
        (50, 49, True),
        (_BUDGET, _BUDGET - 1, True),
        (_BUDGET + 1, _BUDGET - 1, False),
        (_BUDGET * 5, _BUDGET - 1, False),
    ],
    ids=["short", "exactly-the-budget", "one-past-it", "far-past-it"],
)
def test_a_sheet_is_walked_no_further_than_the_budget(
    total: int, rows: int, exact: bool
) -> None:
    """Only a table whose end is seen inside the budget gets an exact count;
    the lookahead row is what separates the two."""
    walked = _CountingRows(total)

    table = spreadsheet_handler_module._read_sheet(walked).table

    assert (table["num_rows"], table["row_count_is_exact"]) == (rows, exact)
    assert walked.seen <= _BUDGET + 1, "the reader walked past its budget"
    assert walked.closed, "a partly consumed row iterator was left open"


def test_a_count_that_is_not_exact_is_described_as_a_lower_bound(
    handler: SpreadsheetHandler, tmp_path: Path
) -> None:
    """A floor printed as a total is the kind of wrong a reader cannot see."""
    path = _book(
        tmp_path / "long.xlsx",
        {"long": [["gene", "padj"]] + [[f"G{n}", 0.01] for n in range(_BUDGET)]},
    )
    meta = handler.extract(make_source(path))
    meta["relative_path"] = "long.xlsx"

    built = handler.build_croissant([meta], ["file_0"])

    assert meta["sheets"][0]["row_count_is_exact"] is False
    assert f"at least {_BUDGET - 1} rows" in built.record_sets[0].description


def test_a_sheet_declaring_the_wrong_size_is_read_by_its_cells(
    handler: SpreadsheetHandler, tmp_path: Path
) -> None:
    """Declared dimensions are not trusted: a writer that understates them
    drops a column and a row without a word.

    https://openpyxl.readthedocs.io/en/stable/optimized.html
    """
    path = _book(tmp_path / "dims.xlsx", {"s": [["a", "b"], [1, 2], [3, 4]]})
    _understate_dimensions(path)

    (sheet,) = _sheets(handler, path)

    assert (sheet["columns"], sheet["num_rows"]) == (["a", "b"], 2)


@pytest.mark.parametrize("state", ["hidden", "veryHidden"])
def test_a_hidden_sheet_is_described_and_said_to_be_hidden(
    handler: SpreadsheetHandler, tmp_path: Path, state: str
) -> None:
    """A reader comparing document against workbook would otherwise find a
    table the document does not mention."""
    path = tmp_path / "plates.xlsx"
    book = openpyxl.Workbook()
    book.remove(book.active)
    for title in ("shown", "tucked"):
        sheet = book.create_sheet(title)
        sheet.append(["a", "b"])
        sheet.append([1, 2])
    book["tucked"].sheet_state = state
    book.save(path)

    meta = handler.extract(make_source(path))
    meta["relative_path"] = "plates.xlsx"
    built = handler.build_croissant([meta], ["file_0"])

    assert [sheet["visibility"] for sheet in meta["sheets"]] == ["visible", state]
    assert state in built.record_sets[1].description
    assert state not in built.record_sets[0].description


def _damage_second_sheet(path: Path, damage) -> None:
    """Rewrite the second worksheet's XML through ``damage``."""
    import zipfile

    with zipfile.ZipFile(io.BytesIO(path.read_bytes())) as archive:
        items = [(i, archive.read(i.filename)) for i in archive.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as out:
        for item, data in items:
            if item.filename == "xl/worksheets/sheet2.xml":
                data = damage(data)
            out.writestr(item, data)


def _two_sheets(path: Path) -> Path:
    return _book(path, {"good": [["a", "b"], [1, 2]], "bad": [["c", "d"], [3, 4]]})


def test_a_sheet_that_fails_while_being_read_costs_only_that_sheet(
    handler: SpreadsheetHandler, tmp_path: Path
) -> None:
    """Real damaged XML rather than an injected exception, so this exercises
    the reader and not just the catch around it."""
    path = _two_sheets(tmp_path / "damaged.xlsx")
    _damage_second_sheet(path, lambda xml: xml.replace(b"<v>3</v>", b"<v>NaN-ish</v>"))

    meta = handler.extract(make_source(path))

    assert [sheet["name"] for sheet in meta["sheets"]] == ["good"]
    assert [(d.code, d.part) for d in meta["diagnostics"]] == [
        (DiagnosticCode.SHEET_SKIPPED, "bad")
    ]
    assert "could not be read" in meta["diagnostics"][0].detail


def test_a_workbook_whose_package_will_not_parse_is_lost_whole(
    handler: SpreadsheetHandler, tmp_path: Path
) -> None:
    """The limit of the isolation above: a worksheet part that is not XML
    fails while the package is opened, and the healthy sheet goes with it."""
    path = _two_sheets(tmp_path / "unopenable.xlsx")
    _damage_second_sheet(path, lambda xml: b"<<<not xml" + xml)

    with pytest.raises(ValueError, match="unopenable.xlsx"):
        handler.extract(make_source(path))


_BROKEN_CUSTOM = (
    b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    b'<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006'
    b'/custom-properties" xmlns:vt="http://schemas.openxmlformats.org/office'
    b'Document/2006/docPropsVTypes"><property fmtid="{D5CDD505-2E9C-101B-9397-'
    b'08002B2CF9AE}" pid="2"><vt:lpwstr>x</vt:lpwstr></property></Properties>'
)


def _with_custom_properties(path: Path, payload: bytes) -> None:
    """Attach a custom-properties part, its relationship and its content type.

    A property carrying no name, as in GSE131907_Lung_Cancer_Feature_Summary.
    """
    import zipfile

    source = path.read_bytes()
    with zipfile.ZipFile(io.BytesIO(source)) as archive:
        items = [(i, archive.read(i.filename)) for i in archive.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as out:
        for item, data in items:
            if item.filename == "[Content_Types].xml":
                data = data.replace(
                    b"</Types>",
                    b'<Override PartName="/docProps/custom.xml" ContentType='
                    b'"application/vnd.openxmlformats-officedocument.custom-'
                    b'properties+xml"/></Types>',
                )
            elif item.filename == "_rels/.rels":
                data = data.replace(
                    b"</Relationships>",
                    b'<Relationship Id="rIdCustom" Type="http://schemas.open'
                    b"xmlformats.org/officeDocument/2006/relationships/custom-"
                    b'properties" Target="docProps/custom.xml"/></Relationships>',
                )
            out.writestr(item, data)
        out.writestr("docProps/custom.xml", payload)


def test_a_workbook_with_malformed_properties_is_read_without_them(
    handler: SpreadsheetHandler, tmp_path: Path
) -> None:
    """Dropped on a copy, only after the part is confirmed unreadable, and
    the identity the document carries stays the file's own."""
    path = _book(tmp_path / "geo.xlsx", {"summary": [["gene", "padj"], ["TP53", 0.01]]})
    _with_custom_properties(path, _BROKEN_CUSTOM)
    # From the bytes before the read: comparing a result against its own
    # source object would hold however the file was mangled in between.
    before = path.read_bytes()
    digest = hashlib.sha256(before).hexdigest()

    meta = handler.extract(make_source(path))

    assert [sheet["columns"] for sheet in meta["sheets"]] == [["gene", "padj"]]
    assert [d.code for d in meta["diagnostics"]] == [DiagnosticCode.PROPERTIES_IGNORED]
    assert path.read_bytes() == before, "the workbook on disk was modified"
    assert (meta["sha256"], meta["file_size"]) == (digest, len(before))
    assert meta["file_name"] == "geo.xlsx"
    assert "custom document properties" in meta["description"]


@pytest.mark.parametrize("suffix", [".xlsx", ".xlsm"])
def test_an_ole2_container_under_a_modern_name_is_told_what_it_is(
    tmp_path: Path, suffix: str
) -> None:
    """Asserted through the pipeline, because the point is the reason that
    reaches the user: claimed, then failed, and never `no handler`."""
    (tmp_path / f"locked{suffix}").write_bytes(
        b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 600
    )
    _book(tmp_path / "readable.xlsx", {"s": [["a", "b"], [1, 2]]})

    _, report = bake_with_report(tmp_path)
    payload = report.to_dict()

    (entry,) = [f for f in payload["files"] if f["path"].startswith("locked")]
    assert entry["reason"] == "extract_failed"
    assert payload["by_reason"] == {"extract_failed": 1}
    assert "no_handler" not in payload["by_reason"]
    assert "encrypted or a legacy workbook" in entry["detail"]
    assert "unencrypted workbook" in entry["detail"]


def test_a_workbook_that_cannot_be_read_fails_by_name(
    handler: SpreadsheetHandler, tmp_path: Path
) -> None:
    """Recovery is for the one part that is known to be optional; anything
    else still fails."""
    path = tmp_path / "book.xlsx"
    path.write_bytes(b"PK\x03\x04" + b"\x00" * 300)

    with pytest.raises(ValueError, match="Failed to read spreadsheet") as raised:
        handler.extract(make_source(path))
    assert "book.xlsx" in str(raised.value)


def test_properties_recovery_does_not_rescue_an_otherwise_broken_workbook(
    handler: SpreadsheetHandler, tmp_path: Path
) -> None:
    """A workbook whose properties are malformed *and* whose sheets are gone is
    still a failure: the retry drops one part, it does not paper over the file.
    """
    import zipfile

    path = _book(tmp_path / "both.xlsx", {"s": [["a"], [1]]})
    _with_custom_properties(path, _BROKEN_CUSTOM)
    source = path.read_bytes()
    with zipfile.ZipFile(io.BytesIO(source)) as archive:
        items = [(i, archive.read(i.filename)) for i in archive.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as out:
        for item, data in items:
            if item.filename.startswith("xl/worksheets/"):
                data = b"<not xml"
            out.writestr(item, data)

    with pytest.raises(ValueError, match="both.xlsx"):
        handler.extract(make_source(path))


@pytest.mark.parametrize("reverse", [False, True], ids=["as-walked", "reversed"])
@pytest.mark.parametrize("workers", [1, 4], ids=["serial", "parallel"])
def test_identifiers_do_not_depend_on_the_order_files_were_found(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, reverse: bool, workers: int
) -> None:
    """`a.xlsx` sheet `b_c` and `a_b.xlsx` sheet `c` both want `a_b_c`. Which
    one moves has to follow from the names, not from the walk order."""
    _book(tmp_path / "a.xlsx", {"b_c": [["v"], [1]]})
    _book(tmp_path / "a_b.xlsx", {"c": [["v"], [2]]})
    if reverse:
        from croissant_baker import scan

        discover = scan.discover_files
        monkeypatch.setattr(
            scan,
            "discover_files",
            lambda *args, **kwargs: list(reversed(discover(*args, **kwargs))),
        )

    document = bake(tmp_path, max_workers=workers)

    assert sorted(node["@id"] for node in record_sets(document)) == [
        "a_b_c",
        "a_b_c__2",
    ]
    # Allocated in path order, so the bare id goes to the file that sorts
    # first and stays there however the walk returned the two.
    by_id = {node["@id"]: node["description"] for node in record_sets(document)}
    assert "of a.xlsx" in by_id["a_b_c"], "the id moved with the discovery order"
    assert "of a_b.xlsx" in by_id["a_b_c__2"]


def test_a_partly_described_workbook_reports_its_sheets_everywhere(
    tmp_path: Path,
) -> None:
    """A sheet is not a file, so the workbook stays described and no coverage
    total moves. What it could not describe has to reach all four places a
    user looks: the summary, --verbose, --report, and the manifest.
    """
    dataset = tmp_path / "in"
    dataset.mkdir()
    out = tmp_path / "out"
    out.mkdir()
    _book(
        dataset / "study.xlsx",
        {
            "good": [["a", "b"], [1, 2]],
            "notes": _SIDE_BY_SIDE,
            "header only": [["gene", "log2FC"]],
            "Sheet2": [],
        },
    )
    (dataset / "plain.csv").write_text("id,name\n1,Ada\n")

    quiet = cli(dataset, out / "quiet.jsonld")
    loud = cli(
        dataset, out / "loud.jsonld", "--verbose", "--report", str(out / "r.json")
    )

    assert quiet.exit_code == 0 and loud.exit_code == 0
    assert "Scanned 2 file(s): 2 described, 0 not described." in quiet.output
    assert "sheet not described: 2" in quiet.output
    assert "notes" not in quiet.output, "the default output named a sheet"
    assert "Tip: re-run with --verbose" in quiet.output
    assert "study.xlsx [notes]" in loud.output
    assert "study.xlsx [header only]" in loud.output
    assert "Sheet2" not in loud.output, "an empty sheet was reported"

    payload = json.loads((out / "r.json").read_text())
    assert (payload["described"], payload["undescribed"]) == (2, 0)
    assert payload["by_reason"] == {}
    assert payload["by_diagnostic"] == {"sheet_skipped": 2}
    by_path = {f["path"]: f for f in payload["files"]}
    assert "diagnostics" not in by_path["plain.csv"]
    assert [(d["code"], d["part"]) for d in by_path["study.xlsx"]["diagnostics"]] == [
        ("sheet_skipped", "notes"),
        ("sheet_skipped", "header only"),
    ]
    details = [d["detail"] for d in by_path["study.xlsx"]["diagnostics"]]
    assert "column 3 is empty throughout" in details[0]
    assert "no data under it" in details[1]

    document = json.loads((out / "loud.jsonld").read_text())
    described = {node["name"]: node for node in file_objects(document)}
    description = described["study.xlsx"]["description"]
    assert "'notes' was not described" in description
    assert "'header only' was not described" in description
    assert "Sheet2" not in description
    assert "description" not in described["plain.csv"]


def test_sheets_take_part_in_foreign_key_detection(tmp_path: Path) -> None:
    """Sheets join in with no code of their own; and where two workbooks each
    hold a `samples` sheet, the pass refuses rather than pick a parent."""
    _book(
        tmp_path / "study.xlsx",
        {
            "samples": [["sample_id", "tissue"], [1, "lung"]],
            "reads": [["read_id", "sample_id"], [9, 1]],
        },
    )
    document = bake(tmp_path, detect_references=True)
    linked = {
        (node["@id"], field["name"]): field["references"]["field"]["@id"]
        for node in record_sets(document)
        for field in node["field"]
        if field.get("references")
    }
    assert linked == {("study_reads", "sample_id"): "study_samples/sample_id"}

    _book(tmp_path / "other.xlsx", {"samples": [["sample_id", "tissue"], [2, "liver"]]})
    document = bake(tmp_path, detect_references=True)

    assert not [
        field
        for node in record_sets(document)
        for field in node["field"]
        if field.get("references")
    ], "an ambiguous parent was linked anyway"


def _cache_formulas(path: Path, cached: dict) -> None:
    """Give cells a formula and the value Excel would have cached for it.

    openpyxl writes formulas without a cache, so this goes into the XML.
    """
    import zipfile

    with zipfile.ZipFile(io.BytesIO(path.read_bytes())) as archive:
        items = [(i, archive.read(i.filename)) for i in archive.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as out:
        for item, data in items:
            if item.filename == "xl/worksheets/sheet1.xml":
                for ref, value in cached.items():
                    before = f'<c r="{ref}" t="n"><v>0</v></c>'.encode()
                    after = f'<c r="{ref}"><f>A1+1</f><v>{value}</v></c>'.encode()
                    assert before in data, f"{ref} is not where the fixture put it"
                    data = data.replace(before, after)
            out.writestr(item, data)


def test_a_formula_is_typed_by_the_value_excel_cached_for_it(
    handler: SpreadsheetHandler, tmp_path: Path
) -> None:
    """Cannot pass by accident: the cached values are fractions, so reading
    them gives cr:Float64, and ignoring them leaves the sheet with no data."""
    path = _book(tmp_path / "cached.xlsx", {"calc": [["total"], [0], [0]]})
    _cache_formulas(path, {"A2": "2.5", "A3": "4.5"})

    (sheet,) = _sheets(handler, path)

    assert (sheet["columns"], sheet["num_rows"]) == (["total"], 2)
    assert sheet["column_types"] == ["cr:Float64"]
